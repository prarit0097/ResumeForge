"""Upload flow: posting a real file to /enhance/ creates a draft and redirects."""
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

pytestmark = pytest.mark.django_db

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _make_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Asha Verma")
    doc.add_paragraph("asha@example.com | +91 98765 43210")
    doc.add_paragraph("Senior Software Engineer who ships reliable systems.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_docx_creates_draft_and_redirects(client):
    upload = SimpleUploadedFile("resume.docx", _make_docx(), content_type=DOCX_MIME)
    resp = client.post("/enhance/", {"resume": upload})
    assert resp.status_code == 302
    # Enhance flow lands on the before/after comparison page first.
    assert "/compare/" in resp["Location"]


def test_upload_with_jd_stores_it_and_compare_shows_match(client):
    from apps.resumes.models import Resume
    upload = SimpleUploadedFile("resume.docx", _make_docx(), content_type=DOCX_MIME)
    resp = client.post("/enhance/", {
        "resume": upload, "jd": "Looking for a backend engineer with Python and Django."})
    assert resp.status_code == 302
    rid = resp["Location"].split("/r/")[1].split("/")[0]
    r = Resume.objects.get(pk=rid)
    assert "Python" in r.job_description  # JD persisted
    # The compare page now renders a JD-match score block.
    page = client.get(resp["Location"]).content.decode()
    assert "Match to your target job" in page


def test_upload_without_jd_has_no_match_block(client):
    upload = SimpleUploadedFile("resume.docx", _make_docx(), content_type=DOCX_MIME)
    resp = client.post("/enhance/", {"resume": upload})
    page = client.get(resp["Location"]).content.decode()
    assert "Match to your target job" not in page


def test_upload_rejects_wrong_type(client):
    bad = SimpleUploadedFile("resume.txt", b"hello", content_type="text/plain")
    resp = client.post("/enhance/", {"resume": bad})
    assert resp.status_code == 200  # re-renders the page with an error
    assert b"PDF or Word" in resp.content


def test_upload_without_file_shows_error(client):
    resp = client.post("/enhance/", {})
    assert resp.status_code == 200
    assert b"choose a file" in resp.content
